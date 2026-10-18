# SPDX-License-Identifier: Apache-2.0
"""Import XLSForm (ARC-27) : classeur `survey`/`choices`/`settings` -> plan de collecte.

Fichier NON fiable : bornes de taille décompressée, de feuilles, de lignes, de colonnes et de
longueur de cellule ; `data_only=True` (jamais de formule évaluée) ; XML lu par openpyxl, qui
délègue à `defusedxml` (installé : pas de XXE). Rien n'est perdu en silence : tout élément non
porté (ou approximé) figure dans `report`."""

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.cell_safety import unsanitize_cell
from app.collections.domains import cel_syntax_error

MAX_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED = 30 * 1024 * 1024
MAX_SHEETS = 20
MAX_SURVEY_ROWS = 1000
MAX_CHOICE_ROWS = 10_000
MAX_COLUMNS = 60
MAX_CELL = 2000
_MAX_NAME = 59


class XlsFormError(Exception):
    """Classeur inexploitable (message affichable tel quel)."""


@dataclass
class ReportItem:
    code: str  # not_ported | approximated | renamed | ignored
    message: str
    row: int | None = None
    name: str | None = None

    def as_dict(self) -> dict:
        return {"code": self.code, "message": self.message, "row": self.row, "name": self.name}


@dataclass
class XlsFormPlan:
    title: str
    columns: list[dict] = field(default_factory=list)  # {name, sqlType}
    geometry_type: str | None = None
    attachment_fields: list[dict] = field(default_factory=list)  # {key, label, kind}
    geometry_required: bool = False
    field_domains: dict[str, dict] = field(default_factory=dict)
    attribute_rules: list[dict] = field(default_factory=list)
    form_fields: list[dict] = field(default_factory=list)
    report: list[ReportItem] = field(default_factory=list)


_SQL = {
    "text": ("text", "string"),
    "integer": ("integer", "integer"),
    "int": ("integer", "integer"),
    "decimal": ("double precision", "number"),
    "date": ("date", "date"),
    "datetime": ("timestamptz", "datetime"),
    "acknowledge": ("boolean", "boolean"),
    "barcode": ("text", "string"),
    "time": ("text", "string"),
}
_APPROX = {
    "time": "heure stockée comme texte",
    "barcode": "code-barres saisi comme texte",
}
_GEO = {"geopoint": "Point", "geotrace": "LineString", "geoshape": "Polygon"}
_FILES = {"image", "photo", "audio", "video", "file"}
_FILE_KIND = {"image": "image", "photo": "image", "audio": "audio", "video": "video"}
_METADATA = {
    "start",
    "end",
    "today",
    "deviceid",
    "subscriberid",
    "simserial",
    "phonenumber",
    "username",
    "email",
    "audit",
}
_RESERVED = {"id", "tenant_id", "geom"}
_UNSUPPORTED_COLS = (
    "appearance",
    "choice_filter",
    "read_only",
    "repeat_count",
    "parameters",
)


# ------------------------------------------------------------------ conversion XPath -> CEL

_TOKEN = re.compile(
    r"""
    (?P<str>'[^']*'|"[^"]*")
  | (?P<ref>\$\{[^}]+\})
  | (?P<fn>string-length(?=\s*\())
  | (?P<word>[A-Za-z_][A-Za-z_0-9]*)
  | (?P<num>\d+(?:\.\d+)?)
  | (?P<dot>(?<![\w.])\.(?![\w.]))
  | (?P<op>!=|<=|>=|=|<|>|\(|\)|,|\+|-|\*)
  | (?P<ws>\s+)
    """,
    re.VERBOSE,
)


class _Unsupported(Exception):
    pass


def _tokenize(expr: str) -> list[tuple[str, str]]:
    pos = 0
    toks: list[tuple[str, str]] = []
    while pos < len(expr):
        m = _TOKEN.match(expr, pos)
        if not m:
            raise _Unsupported(f"caractère non pris en charge : {expr[pos]!r}")
        pos = m.end()
        if m.lastgroup != "ws":
            toks.append((m.lastgroup or "", m.group()))
    return toks


def _call_args(
    toks: list[tuple[str, str]], open_idx: int
) -> tuple[list[list[tuple[str, str]]], int]:
    """Arguments (listes de jetons) de l'appel dont la parenthèse ouvrante est à `open_idx`,
    et l'index de la parenthèse fermante."""
    depth = 0
    args: list[list[tuple[str, str]]] = [[]]
    for j in range(open_idx, len(toks)):
        tok = toks[j][1]
        if tok == "(":
            depth += 1
            if depth == 1:
                continue
        elif tok == ")":
            depth -= 1
            if depth == 0:
                return ([a for a in args if a] if args != [[]] else []), j
        elif tok == "," and depth == 1:
            args.append([])
            continue
        args[-1].append(toks[j])
    raise _Unsupported("parenthèse non fermée")


def _translate(toks: list[tuple[str, str]], names: dict[str, str], current: str | None) -> str:
    out: list[str] = []
    i = 0
    while i < len(toks):
        kind, tok = toks[i]
        nxt = toks[i + 1][1] if i + 1 < len(toks) else ""
        if kind == "ref":
            name = tok[2:-1].strip()
            if name not in names:
                raise _Unsupported(f"champ inconnu : {name}")
            out.append(names[name])
        elif kind == "dot":
            if current is None:
                raise _Unsupported("'.' hors contrainte")
            out.append(current)
        elif kind == "fn":  # string-length(x) -> size(x) (identique pour une valeur présente)
            args, close = _call_args(toks, i + 1)
            if len(args) != 1:
                raise _Unsupported("string-length() attend 1 argument")
            out.append(f"size({_translate(args[0], names, current)})")
            i = close
        elif kind == "word":
            low = tok.lower()
            if low in ("and", "or"):
                out.append("&&" if low == "and" else "||")
            elif low == "not" and nxt == "(":
                out.append("!")
            elif low in ("true", "false") and nxt == "(":
                # true() / false() : on saute le couple de parenthèses
                if toks[i + 2][1] != ")":
                    raise _Unsupported(f"{tok}() avec arguments")
                out.append(low)
                i += 2
            elif low == "selected" and nxt == "(":
                # selected(${f}, 'v') ; valide seulement pour select_one (valeur unique)
                try:
                    ref, comma, lit, close = (toks[i + 2], toks[i + 3], toks[i + 4], toks[i + 5])
                except IndexError as exc:
                    raise _Unsupported("selected() mal formé") from exc
                if ref[0] != "ref" or comma[1] != "," or lit[0] != "str" or close[1] != ")":
                    raise _Unsupported("selected() mal formé")
                name = ref[1][2:-1].strip()
                if name not in names:
                    raise _Unsupported(f"champ inconnu : {name}")
                out.append(f"({names[name]} == {lit[1]})")
                i += 5
            elif low == "if" and nxt == "(":
                args, close = _call_args(toks, i + 1)
                if len(args) != 3:
                    raise _Unsupported("if() attend 3 arguments")
                c, a, b = (_translate(x, names, current) for x in args)
                out.append(f"(({c}) ? ({a}) : ({b}))")
                i = close
            elif low == "coalesce" and nxt == "(":
                args, close = _call_args(toks, i + 1)
                if len(args) < 2:
                    raise _Unsupported("coalesce() attend au moins 2 arguments")
                parts = [_translate(x, names, current) for x in args]
                expr = parts[-1]
                for p in reversed(parts[:-1]):  # coalesce(a,b,c) = a ?? (b ?? c), a évalué 2 fois
                    expr = f"(({p}) != null ? ({p}) : ({expr}))"
                out.append(expr)
                i = close
            elif low == "regex" and nxt == "(":
                args, close = _call_args(toks, i + 1)
                if len(args) != 2:
                    raise _Unsupported("regex() attend 2 arguments")
                v, pat = (_translate(x, names, current) for x in args)
                out.append(f"({v}).matches({pat})")
                i = close
            elif low == "today" and nxt == "(":
                raise _Unsupported("today() n'a pas d'équivalent CEL (pas d'horloge côté règle)")
            elif low == "mod":
                out.append("%")
            elif low == "div":
                raise _Unsupported("opérateur div")
            elif nxt == "(":
                raise _Unsupported(f"fonction {tok}()")
            else:
                raise _Unsupported(f"mot non pris en charge : {tok}")
        elif kind == "op":
            out.append("==" if tok == "=" else tok)
        else:  # str, num
            out.append(tok)
        i += 1
    return " ".join(out)


def to_cel(expr: str, names: dict[str, str], *, current: str | None = None) -> str:
    """Traduit le sous-ensemble utile d'une expression XPath XLSForm en CEL (champs = noms de
    colonnes). Lève _Unsupported pour tout le reste : jamais d'approximation silencieuse.
    Fonctions : not, true/false, selected (select_one), if, coalesce, regex, string-length ;
    `today()` est refusée (aucune horloge dans l'évaluation des règles)."""
    cel = _translate(_tokenize(expr), names, current)
    cel = cel.replace("! (", "!(").replace("( ", "(").replace(" )", ")")
    if not cel.strip():
        raise _Unsupported("expression vide")
    if err := cel_syntax_error(cel):
        raise _Unsupported(f"CEL invalide ({err})")
    return cel


def _infer_sql_type(expr: str, names: dict[str, str], col_types: dict[str, str]) -> str:
    """Type SQL du résultat d'un `calculate`, déduit de la forme de l'expression ; lève
    _Unsupported quand il est indéterminé (aucune supposition silencieuse)."""
    toks = _tokenize(expr)
    kinds = {k for k, _ in toks}
    words = {t.lower() for k, t in toks if k == "word"}
    if words & {"if", "coalesce", "regex", "selected", "and", "or", "not"} or any(
        t in ("=", "!=", "<", ">", "<=", ">=") for k, t in toks if k == "op"
    ):
        raise _Unsupported("type du résultat indéterminé (condition)")
    refs = [names[t[2:-1].strip()] for k, t in toks if k == "ref"]
    ref_types = {col_types.get(r) for r in refs}
    ops = {t for k, t in toks if k == "op"} & {"+", "-", "*", "%"}
    if "fn" in kinds:
        return "double precision"
    if "str" in kinds:
        return "text"
    if ops and ref_types <= {"double precision", "integer"}:
        return "double precision"
    if not ops and len(toks) == 1 and len(ref_types) == 1 and None not in ref_types:
        return next(iter(ref_types))  # simple recopie d'un autre champ
    raise _Unsupported("type du résultat indéterminé")


def _cel_literal(value: str, sql_type: str, *, multiple: bool = False) -> str:
    """Valeur littérale de `default` -> littéral CEL du type de la colonne ; _Unsupported sinon."""
    if multiple:
        return "[" + ", ".join(_cel_literal(v, "text") for v in value.split()) + "]"
    if sql_type in ("text", "date", "timestamptz"):
        return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"
    if sql_type == "integer":
        try:
            return str(int(value))
        except ValueError as exc:
            raise _Unsupported(f"valeur par défaut non entière : {value!r}") from exc
    if sql_type == "double precision":
        try:
            return repr(float(value))
        except ValueError as exc:
            raise _Unsupported(f"valeur par défaut non numérique : {value!r}") from exc
    raise _Unsupported(f"valeur par défaut non portée pour le type {sql_type}")


# ------------------------------------------------------------------ lecture du classeur


def _cell(value: Any, row: int, col: str) -> str:
    if value is None:
        return ""
    s = str(unsanitize_cell(value) if isinstance(value, str) else value).strip()
    if len(s) > MAX_CELL:
        raise XlsFormError(f"cellule trop longue (ligne {row}, colonne {col})")
    return s


def _read_sheet(wb, name: str, max_rows: int) -> list[tuple[int, dict[str, str]]]:
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    cols = [str(h).strip().lower() if h is not None else "" for h in header[:MAX_COLUMNS]]
    if len(header) > MAX_COLUMNS and any(h is not None for h in header[MAX_COLUMNS:]):
        raise XlsFormError(f"feuille {name} : plus de {MAX_COLUMNS} colonnes")
    out: list[tuple[int, dict[str, str]]] = []
    for idx, raw in enumerate(rows, start=2):
        if idx - 1 > max_rows:
            raise XlsFormError(f"feuille {name} : plus de {max_rows} lignes")
        rec = {c: _cell(v, idx, c) for c, v in zip(cols, raw, strict=False) if c}
        if any(rec.values()):
            out.append((idx, rec))
    return out


def _label(rec: dict[str, str]) -> str:
    if rec.get("label"):
        return rec["label"]
    # AUD-115 : choix déterministe (indépendant de l'ordre des colonnes) — français d'abord
    # (langue source du produit), puis ordre alphabétique de la langue.
    translated = sorted(
        ((k, v) for k, v in rec.items() if k.startswith("label::") and v),
        key=lambda kv: (not kv[0].startswith(("label::fr", "label::fran")), kv[0]),
    )
    return translated[0][1] if translated else ""


def _sanitize_name(raw: str) -> str:
    n = re.sub(r"[^A-Za-z0-9_]", "_", raw)
    if not n or n[0].isdigit():
        n = "q_" + n
    if n in _RESERVED:
        n += "_q"
    return n[:_MAX_NAME]


def _truthy(v: str) -> bool:
    return v.strip().lower() in {"yes", "true", "true()", "1"}


def parse_xlsform(content: bytes) -> XlsFormPlan:
    if len(content) > MAX_BYTES:
        raise XlsFormError("fichier trop volumineux")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise XlsFormError("classeur trop volumineux une fois décompressé")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise XlsFormError("fichier XLSX illisible") from exc
    if len(wb.sheetnames) > MAX_SHEETS:
        raise XlsFormError(f"plus de {MAX_SHEETS} feuilles")
    if "survey" not in wb.sheetnames:
        raise XlsFormError("feuille « survey » absente : ce n'est pas un XLSForm")

    survey = _read_sheet(wb, "survey", MAX_SURVEY_ROWS)
    choices_rows = _read_sheet(wb, "choices", MAX_CHOICE_ROWS)
    settings = _read_sheet(wb, "settings", 5)
    choices: dict[str, list[tuple[str, str]]] = {}
    for _, c in choices_rows:
        if c.get("list_name") and c.get("name"):
            choices.setdefault(c["list_name"], []).append((c["name"], _label(c) or c["name"]))

    title = (settings[0][1].get("form_title") if settings else "") or "Formulaire XLSForm"
    plan = XlsFormPlan(title=title)
    rep = plan.report

    # 1re passe : noms de colonnes finaux (les expressions y font référence).
    names: dict[str, str] = {}
    used: set[str] = set()
    col_types: dict[str, str] = {}  # nom final -> type SQL (pour typer un `calculate`)
    calcs: dict[str, tuple[str, str]] = {}  # nom brut -> (CEL, type SQL) des calculate portables
    depth0 = 0
    for _, q in survey:
        raw = q.get("name", "")
        parts0 = q.get("type", "").split()
        typ = parts0[0].lower() if parts0 else ""
        if typ in ("begin_repeat", "begin repeat"):
            depth0 += 1
        elif typ in ("end_repeat", "end repeat"):
            depth0 = max(0, depth0 - 1)
        if depth0 or not raw or typ in ("end", "end_group", "end_repeat", "begin", "note"):
            continue
        if typ == "calculate":
            # nommable seulement si portable (sinon une référence pointerait une colonne absente)
            try:
                expr = to_cel(q.get("calculation", ""), names)
                sql = _infer_sql_type(q.get("calculation", ""), names, col_types)
            except _Unsupported:
                continue
        n = _sanitize_name(raw)
        base, k = n, 2
        while n in used:
            n = f"{base[: _MAX_NAME - 3]}_{k}"
            k += 1
        used.add(n)
        names[raw] = n
        if typ == "calculate":
            calcs[raw] = (expr, sql)
            col_types[n] = sql
        elif typ in _SQL:
            col_types[n] = _SQL[typ][0]
        elif typ in ("select_one", "select_multiple", "rank"):
            col_types[n] = "text[]" if typ == "select_multiple" else "text"

    depth_repeat = 0
    order = 0
    for row, q in survey:
        parts = q.get("type", "").split()
        typ = parts[0].lower() if parts else ""
        raw = q.get("name", "")
        if typ in ("begin_repeat", "begin repeat"):
            depth_repeat += 1
            rep.append(ReportItem("not_ported", "répétition (repeat) non portée", row, raw))
            continue
        if typ in ("end_repeat", "end repeat"):
            depth_repeat = max(0, depth_repeat - 1)
            continue
        if depth_repeat:
            rep.append(
                ReportItem("not_ported", "question dans une répétition non portée", row, raw)
            )
            continue
        if typ in ("begin_group", "begin group"):
            rep.append(ReportItem("approximated", "groupe aplati (ordre conservé)", row, raw))
            continue
        if typ in ("end_group", "end group", ""):
            continue
        if typ in _METADATA:
            rep.append(ReportItem("ignored", f"métadonnée {typ} non portée", row, raw))
            continue
        if typ == "note":
            rep.append(ReportItem("not_ported", "note (affichage seul) non portée", row, raw))
            continue
        if typ == "calculate":
            if raw in calcs:
                cel, sql = calcs[raw]
                plan.columns.append({"name": names[raw], "sqlType": sql})
                plan.attribute_rules.append(
                    {
                        "id": f"xf-calc-{names[raw].lower()}"[:64],
                        "kind": "calculation",
                        "field": names[raw],
                        "expression": cel,
                        "mode": "derived",
                    }
                )
                rep.append(
                    ReportItem(
                        "approximated",
                        f"calculate porté en règle serveur (colonne {sql} déduite de l'expression)",
                        row,
                        raw,
                    )
                )
            else:
                rep.append(
                    ReportItem(
                        "not_ported",
                        "calculate non porté (expression ou type du résultat non traduisible)",
                        row,
                        raw,
                    )
                )
            continue
        if not raw or raw not in names:
            rep.append(ReportItem("not_ported", "question sans nom", row, raw or None))
            continue
        name = names[raw]
        if name != raw:
            rep.append(ReportItem("renamed", f"renommé en {name}", row, raw))
        label = _label(q) or raw
        for k in q:
            if k.startswith("label::") and q[k] and q.get("label"):
                rep.append(ReportItem("ignored", f"traduction {k} ignorée", row, raw))
        for c in _UNSUPPORTED_COLS:
            if q.get(c):
                rep.append(ReportItem("not_ported", f"colonne {c} non portée", row, raw))

        entry: dict[str, Any] = {
            "name": name,
            "label": label,
            "order": order,
            "hidden": False,
            "required": _truthy(q.get("required", "")),
        }
        if q.get("hint"):
            entry["hint"] = q["hint"]

        if typ in _GEO:
            if plan.geometry_type is None:
                plan.geometry_type = _GEO[typ]
                if _truthy(q.get("required", "")):
                    if q.get("relevant"):
                        rep.append(
                            ReportItem(
                                "approximated",
                                "géométrie requise non imposée (question conditionnelle)",
                                row,
                                raw,
                            )
                        )
                    else:
                        plan.geometry_required = True
            else:
                rep.append(
                    ReportItem("not_ported", "2e géométrie : une seule par collection", row, raw)
                )
            order += 1
            continue
        if typ in _FILES:
            plan.attachment_fields.append(
                {"key": name, "label": label, "kind": _FILE_KIND.get(typ, "file")}
            )
            entry["type"] = "attachment"
            entry["kind"] = _FILE_KIND.get(typ, "file")
            if entry["required"]:
                rep.append(
                    ReportItem(
                        "approximated",
                        "pièce jointe requise : contrôlée dans le formulaire seulement "
                        "(le fichier arrive après la saisie)",
                        row,
                        raw,
                    )
                )
        elif typ in ("select_one", "select_multiple", "rank"):
            lst = parts[1] if len(parts) > 1 else ""
            values = choices.get(lst)
            if not values:
                rep.append(
                    ReportItem("not_ported", f"liste de choix {lst!r} absente ou vide", row, raw)
                )
                continue
            vals = [v for v, _ in values]
            dom: dict[str, Any] = {"type": "list", "values": vals}
            if any(lbl != v for v, lbl in values):
                dom["labels"] = {v: lbl for v, lbl in values}
            plan.field_domains[name] = dom
            entry["type"] = "string"
            if typ == "select_one":
                plan.columns.append({"name": name, "sqlType": "text"})
            else:
                plan.columns.append({"name": name, "sqlType": "text[]"})
                entry["multiple"] = True
        elif typ in _SQL:
            sql, ftype = _SQL[typ]
            plan.columns.append({"name": name, "sqlType": sql})
            entry["type"] = ftype
            if typ in _APPROX:
                rep.append(ReportItem("approximated", _APPROX[typ], row, raw))
        else:
            rep.append(ReportItem("not_ported", f"type {typ!r} non pris en charge", row, raw))
            continue

        if q.get("relevant"):
            try:
                entry["visibleWhen"] = to_cel(q["relevant"], names)
            except _Unsupported as exc:
                rep.append(ReportItem("not_ported", f"relevant non porté ({exc})", row, raw))
        if q.get("default") and entry["type"] != "attachment":
            dflt = q["default"]
            try:
                if "${" in dflt or "(" in dflt:
                    expr = to_cel(dflt, names)
                else:
                    expr = _cel_literal(
                        dflt, col_types.get(name, "text"), multiple=bool(entry.get("multiple"))
                    )
                plan.attribute_rules.append(
                    {
                        "id": f"xf-default-{name.lower()}"[:64],
                        "kind": "calculation",
                        "field": name,
                        "expression": expr,
                        "mode": "default",
                    }
                )
            except _Unsupported as exc:
                rep.append(ReportItem("not_ported", f"default non porté ({exc})", row, raw))
        if entry["required"] and entry["type"] != "attachment":
            col_t = col_types.get(name, "text")
            if col_t == "text":
                present = f"{name} != null && {name} != ''"
            elif col_t.endswith("[]"):
                present = f"{name} != null && size({name}) > 0"
            else:
                present = f"{name} != null"
            if q.get("relevant") and "visibleWhen" not in entry:
                rep.append(
                    ReportItem(
                        "approximated",
                        "required non imposé côté serveur : relevant non porté "
                        "(le champ peut être masqué)",
                        row,
                        raw,
                    )
                )
            else:
                if q.get("relevant"):
                    present = f"!({entry['visibleWhen']}) || ({present})"
                plan.attribute_rules.append(
                    {
                        "id": f"xf-required-{name.lower()}"[:64],
                        "kind": "constraint",
                        "field": name,
                        "expression": present,
                        "message": f"« {label} » est obligatoire",
                    }
                )
        if q.get("constraint") and entry["type"] != "attachment":
            try:
                expr = to_cel(q["constraint"], names, current=name)
                plan.attribute_rules.append(
                    {
                        "id": f"xf-{len(plan.attribute_rules) + 1}-{name.lower()}"[:64],
                        "kind": "constraint",
                        "field": name,
                        "expression": expr,
                        "message": q.get("constraint_message") or "Valeur invalide",
                    }
                )
            except _Unsupported as exc:
                rep.append(ReportItem("not_ported", f"constraint non portée ({exc})", row, raw))
        plan.form_fields.append(entry)
        order += 1

    if not plan.form_fields and plan.geometry_type is None:
        raise XlsFormError("aucune question exploitable dans la feuille survey")
    return plan
