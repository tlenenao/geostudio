# SPDX-License-Identifier: Apache-2.0
"""P13 (RC-7) : fuites de lecture et surfaces publiques."""

import io

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app import db
from app.analytics.export import rows_to_csv, rows_to_xlsx
from app.auth.dependency import get_current_user
from app.collections.routes import _is_system_table
from app.db import init_db, make_engine, make_session_factory, request_scoped_session
from app.extensions.schemas import ExtensionCreate, ExtensionPatch
from app.items import repository as items_repo
from app.main import create_app
from app.tenants.repository import get_or_create_default_tenant
from app.users.repository import get_or_create_user


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setenv("CORE_ETL_ENABLED", "true")
    engine = make_engine("sqlite+pysqlite:///:memory:")
    init_db(engine)
    Session = make_session_factory(engine)
    with Session() as s:
        tenant = get_or_create_default_tenant(s)
        user = get_or_create_user(
            s,
            tenant_id=tenant.id,
            oidc_sub="sub-1",
            username="alice",
            email=None,
            first_name="",
            last_name="",
        )
        s.commit()
    app = create_app()

    def override_session():
        with request_scoped_session(Session) as session:
            yield session

    app.dependency_overrides[db.get_session] = override_session
    app.dependency_overrides[get_current_user] = lambda: user
    c = TestClient(app)
    c.app_ = app  # type: ignore[attr-defined]
    c.session_factory = Session  # type: ignore[attr-defined]
    c.user = user  # type: ignore[attr-defined]
    c.tenant = tenant  # type: ignore[attr-defined]
    yield c
    engine.dispose()


def _make_reader(client, user=None):
    user = user or client.user
    from app.roles.repository import ensure_built_in_roles
    from app.users.repository import set_user_role

    with client.session_factory() as session:
        roles = ensure_built_in_roles(session, tenant_id=client.tenant.id)
        set_user_role(
            session,
            tenant_id=client.tenant.id,
            user_id=user.id,
            role_id=roles["reader"].id,
            role_slug="reader",
        )
        session.commit()
    user.role_id = roles["reader"].id


def _seed(client, kind, *, published=False, public=False):
    with client.session_factory() as session:
        item = items_repo.create_item(
            session,
            tenant_id=client.tenant.id,
            owner_id=client.user.id,
            resource_type=kind,
            title=f"t-{kind}",
        )
        item.is_published = published
        item.is_public = public
        session.commit()
        return item.id


# P13.01
@pytest.mark.parametrize("path", ["/v1/items", "/v1/items/facets"])
def test_unknown_scope_is_422(client, path):
    assert client.get(path, params={"scope": "bogus"}).status_code == 422


# P13.02 / P13.05
def test_public_surface_only_serves_public_kinds(client):
    alert_id = _seed(client, "alert", published=True)
    site_id = _seed(client, "site", published=True)
    assert client.get(f"/v1/public/items/{alert_id}").status_code == 404
    assert client.get(f"/v1/public/configs/by-item/{alert_id}").status_code == 404
    assert client.get(f"/v1/public/items/{site_id}").status_code == 200
    kinds = {i["resourceType"] for i in client.get("/v1/public/items").json()["items"]}
    assert kinds == {"site"}


# P13.03
def test_system_tables_are_not_registrable():
    for name in ("user_entity", "credential", "realm", "procrastinate_jobs", "client"):
        assert _is_system_table(name)
    assert not _is_system_table("parcelles")


# P13.04
def test_exports_neutralize_formulas():
    rows = [{"=A": "=1+1", "n": "+cmd", "ok": "x", "num": -3}]
    csv = rows_to_csv(rows).decode()
    assert "'=A" in csv and "'=1+1" in csv and "'+cmd" in csv and "-3" in csv
    ws = load_workbook(io.BytesIO(rows_to_xlsx(rows))).active
    assert [c.value for c in ws[2]] == ["'=1+1", "'+cmd", "x", -3]
    assert ws["A1"].value == "'=A"


# P13.07
def test_pipelines_ops_requires_auth(client):
    def deny():
        raise HTTPException(status_code=401, detail="no session")

    client.app_.dependency_overrides[get_current_user] = deny
    assert client.get("/v1/pipelines/ops").status_code == 401


# P13.08
def test_extension_schema_validation():
    ok = {
        "id": "e",
        "tag": "my-widget",
        "label": "L",
        "moduleUrl": "https://x.test/m.js",
        "defaultSize": {"w": 2, "h": 2},
    }
    ExtensionCreate(**ok)
    for bad in (
        {"moduleUrl": "javascript:alert(1)"},
        {"moduleUrl": "http://x.test/m.js"},
        {"tag": "div"},
        {"tag": "My Widget"},
        {"defaultSize": {"w": 0, "h": 2}},
    ):
        with pytest.raises(ValidationError):
            ExtensionCreate(**{**ok, **bad})
    for bad in ({"label": ""}, {"tag": ""}, {"moduleUrl": ""}, {"moduleUrl": "ftp://x"}):
        with pytest.raises(ValidationError):
            ExtensionPatch(**bad)


# P13.09 / P13.10
def test_reader_cannot_read_acl_nor_group_list(client):
    item_id = _seed(client, "app", public=True)
    # le propriétaire garde l'accès
    assert client.get(f"/v1/items/{item_id}/sharing").status_code == 200
    assert client.get("/v1/groups").status_code == 200
    # un autre utilisateur : lit l'item (public) mais pas son ACL
    with client.session_factory() as s:
        bob = get_or_create_user(
            s,
            tenant_id=client.tenant.id,
            oidc_sub="sub-2",
            username="bob",
            email=None,
            first_name="",
            last_name="",
        )
        s.commit()
    _make_reader(client, bob)
    client.app_.dependency_overrides[get_current_user] = lambda: bob
    assert client.get(f"/v1/items/{item_id}").status_code == 200
    assert client.get(f"/v1/items/{item_id}/sharing").status_code == 403
    assert client.get("/v1/groups").status_code == 403


# P13.11
def test_presign_requires_data_manage(client):
    from app.ingestion import routes as ingestion_routes

    _make_reader(client)
    client.app_.dependency_overrides[ingestion_routes.get_s3_client] = lambda: object()
    client.app_.dependency_overrides[ingestion_routes.get_uploads_bucket] = lambda: "b"
    r = client.post("/v1/uploads/presign", json={"filename": "a.csv", "contentType": "text/csv"})
    assert r.status_code == 403
