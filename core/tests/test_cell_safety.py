import csv
import io

import pytest
from openpyxl import load_workbook

from app.analytics.export import rows_to_csv, rows_to_xlsx
from app.cell_safety import sanitize_cell, unsanitize_cell
from app.ingestion.parsers import GeometryMode, parse_csv_latlon, parse_xlsx_sheet

HOSTILE = ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)", "\t=x", "\r=x"]


@pytest.mark.parametrize("v", HOSTILE)
def test_csv_and_xlsx_prefix_hostile_text(v):
    rows = [{"c": v, "n": -5, "t": "-5", "ok": "x"}]
    got = next(csv.DictReader(io.StringIO(rows_to_csv(rows).decode(), newline="")))
    assert got["c"] == "'" + v and got["n"] == "-5" and got["t"] == "'-5" and got["ok"] == "x"
    ws = load_workbook(io.BytesIO(rows_to_xlsx(rows))).active
    # XML normalise \r en \n à la relecture : artefact du lecteur, le préfixe est bien posé.
    assert [c.value for c in ws[2]] == [("'" + v).replace("\r", "\n"), -5, "'-5", "x"]
    assert ws["A2"].data_type == "s"  # jamais une formule


def test_round_trip_csv_and_xlsx():
    rows = [{"c": "=1+1", "d": "-5", "e": "plain", "f": "'quoted"}]
    mode = GeometryMode(kind="none")
    [(_, props)] = parse_csv_latlon(rows_to_csv(rows), mode)
    assert props == rows[0]
    [(_, props)] = parse_xlsx_sheet(rows_to_xlsx(rows), None, mode)
    assert props == rows[0]


def test_sanitize_inverse():
    assert sanitize_cell(-5) == -5 and sanitize_cell(None) is None
    assert unsanitize_cell("'=a") == "=a" and unsanitize_cell("'abc") == "'abc"
