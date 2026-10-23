# SPDX-License-Identifier: Apache-2.0
"""ARC-27 : parsing XLSForm — fidélité (30 questions), rapport sans perte silencieuse, bornes."""

import io
import zipfile

import pytest
from openpyxl import Workbook

from app.xlsform import parser
from app.xlsform.parser import XlsFormError, parse_xlsform, to_cel


def _book(survey, choices=(), settings=None, extra_sheets=0):
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    for r in survey:
        ws.append(list(r))
    if choices:
        c = wb.create_sheet("choices")
        for r in choices:
            c.append(list(r))
    if settings:
        s = wb.create_sheet("settings")
        for r in settings:
            s.append(list(r))
    for i in range(extra_sheets):
        wb.create_sheet(f"x{i}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEAD = ("type", "name", "label", "hint", "required", "relevant", "constraint", "constraint_message")


def _thirty():
    rows = [HEAD]
    rows.append(("text", "nom", "Nom", "Votre nom", "yes", None, None, None))
    rows.append(
        ("integer", "age", "Âge", None, "yes", None, ". >= 0 and . <= 120", "Âge entre 0 et 120")
    )
    rows.append(("decimal", "surface", "Surface", None, None, None, ". > 0", "Surface positive"))
    rows.append(("date", "jour", "Date", None, None, None, None, None))
    rows.append(("datetime", "horodatage", "Horodatage", None, None, None, None, None))
    rows.append(("select_one etat", "etat", "État", None, "yes", None, None, None))
    rows.append(
        ("text", "motif", "Motif de dégradation", None, None, "${etat} = 'degrade'", None, None)
    )
    rows.append(
        (
            "integer",
            "gravite",
            "Gravité",
            None,
            None,
            "${etat} != 'bon' and ${age} > 18",
            None,
            None,
        )
    )
    rows.append(("select_multiple usages", "usages", "Usages", None, None, None, None, None))
    rows.append(("geopoint", "position", "Position", None, "yes", None, None, None))
    rows.append(("image", "photo", "Photo", None, None, None, None, None))
    rows.append(("acknowledge", "ok", "Je confirme", None, None, None, None, None))
    rows.append(("time", "heure", "Heure", None, None, None, None, None))
    for i in range(14, 31):
        rows.append(("text", f"champ_{i}", f"Champ {i}", None, None, None, None, None))
    return _book(
        rows,
        choices=[
            ("list_name", "name", "label"),
            ("etat", "bon", "Bon"),
            ("etat", "degrade", "Dégradé"),
            ("etat", "hs", "Hors service"),
            ("usages", "a", "A"),
            ("usages", "b", "B"),
        ],
        settings=[("form_title", "form_id"), ("Inspection des bancs", "bancs")],
    )


def test_thirty_questions_import_without_loss():
    plan = parse_xlsform(_thirty())
    assert plan.title == "Inspection des bancs"
    assert plan.geometry_type == "Point"
    by = {f["name"]: f for f in plan.form_fields}
    # 30 questions : 29 champs de formulaire + la géopoint (devenue géométrie de la collection)
    assert len(plan.form_fields) + 1 == 30
    types = {c["name"]: c["sqlType"] for c in plan.columns}
    assert types["nom"] == "text"
    assert types["age"] == "integer"
    assert types["surface"] == "double precision"
    assert types["jour"] == "date"
    assert types["horodatage"] == "timestamptz"
    assert types["ok"] == "boolean"
    assert by["photo"]["type"] == "attachment"
    assert plan.attachment_fields == [{"key": "photo", "label": "Photo", "kind": "image"}]
    # liste de choix -> domaine
    assert plan.field_domains["etat"] == {
        "type": "list",
        "values": ["bon", "degrade", "hs"],
        "labels": {"bon": "Bon", "degrade": "Dégradé", "hs": "Hors service"},
    }
    assert types["usages"] == "text[]"  # select_multiple : tableau, plus de texte libre
    assert plan.geometry_required is True  # géopoint required, plus « non imposée »
    # relevant -> CEL
    assert by["motif"]["visibleWhen"] == "etat == 'degrade'"
    assert by["gravite"]["visibleWhen"] == "etat != 'bon' && age > 18"
    # required, hint, ordre
    assert by["nom"]["required"] is True and by["nom"]["hint"] == "Votre nom"
    assert [f["order"] for f in plan.form_fields] == sorted(f["order"] for f in plan.form_fields)
    # constraint -> règles ARC-14
    rules = {r["field"]: r for r in plan.attribute_rules if not r["id"].startswith("xf-req")}
    assert rules["age"]["expression"] == "age >= 0 && age <= 120"
    assert rules["age"]["message"] == "Âge entre 0 et 120"
    assert rules["surface"]["kind"] == "constraint"
    # rien de perdu en silence : les approximations sont déclarées
    msgs = " | ".join(r.message for r in plan.report)
    assert "heure stockée comme texte" in msgs
    assert "choix multiple" not in msgs and "géométrie requise" not in msgs
    required = {r["field"]: r for r in plan.attribute_rules if r["id"].startswith("xf-req")}
    assert set(required) == {"nom", "age", "etat"}
    assert required["nom"]["expression"] == "nom != null && nom != ''"
    assert required["age"]["expression"] == "age != null"
    assert not [r for r in plan.report if r.code == "not_ported"]


def test_unsupported_items_are_reported_never_silently_dropped():
    survey = [
        HEAD + ("default", "calculation"),
        (
            "text",
            "a",
            "A",
            None,
            None,
            "today() > ${b}",
            "pulldata(., 1) = 2",
            None,
            "today()",
            None,
        ),
        ("text", "b", "B", None, None, None, None, None, None, None),
        ("calculate", "c", "C", None, None, None, None, None, None, "today()"),
        ("note", "n", "Note", None, None, None, None, None, None, None),
        ("begin_repeat", "r", "R", None, None, None, None, None, None, None),
        ("text", "in_rep", "Dans repeat", None, None, None, None, None, None, None),
        ("end_repeat", None, None, None, None, None, None, None, None, None),
        ("geopoint", "g1", "G1", None, None, None, None, None, None, None),
        ("geopoint", "g2", "G2", None, None, None, None, None, None, None),
        ("hologram", "h", "H", None, None, None, None, None, None, None),
    ]
    plan = parse_xlsform(_book(survey))
    rep = {(r.name, r.code): r.message for r in plan.report}
    assert "relevant non porté" in rep[("a", "not_ported")] or any(
        r.name == "a" and "relevant" in r.message for r in plan.report
    )
    assert any(r.name == "a" and "constraint" in r.message for r in plan.report)
    assert any(r.name == "a" and "default" in r.message for r in plan.report)
    for n in ("c", "n", "r", "in_rep", "g2", "h"):
        assert any(r.name == n for r in plan.report), n
    assert "visibleWhen" not in next(f for f in plan.form_fields if f["name"] == "a")
    assert plan.attribute_rules == []


def test_names_are_sanitized_and_reserved_names_renamed():
    survey = [HEAD, ("text", "id", "Id", *[None] * 5), ("text", "a-b.c", "ABC", *[None] * 5)]
    plan = parse_xlsform(_book(survey))
    names = [c["name"] for c in plan.columns]
    assert names == ["id_q", "a_b_c"]
    assert sum(r.code == "renamed" for r in plan.report) == 2


@pytest.mark.parametrize(
    "xpath,cel",
    [
        ("${a} = 'x' or ${b} = 'y'", "a == 'x' || b == 'y'"),
        ("not(${a} = 'x')", "!(a == 'x')"),
        ("selected(${a}, 'x')", "(a == 'x')"),
        ("${n} >= 0.5", "n >= 0.5"),
        ("${a} = 'a = b and c'", "a == 'a = b and c'"),
    ],
)
def test_to_cel(xpath, cel):
    names = {"a": "a", "b": "b", "n": "n"}
    assert to_cel(xpath, names) == cel


def test_to_cel_rejects_unknown_references_and_functions():
    for bad in ("${zz} = 1", "today() = 1", "${a} div 2 = 1"):
        with pytest.raises(Exception):  # noqa: B017 - _Unsupported est interne
            to_cel(bad, {"a": "a"})


# --- fichier non fiable ---------------------------------------------------------------------


def test_not_an_xlsform_and_garbage_are_clean_errors():
    with pytest.raises(XlsFormError):
        parse_xlsform(b"pas un zip")
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(XlsFormError, match="survey"):
        parse_xlsform(buf.getvalue())


def test_bounds(monkeypatch):
    with pytest.raises(XlsFormError, match="volumineux"):
        parse_xlsform(b"x" * (parser.MAX_BYTES + 1))
    monkeypatch.setattr(parser, "MAX_SURVEY_ROWS", 3)
    rows = [HEAD] + [("text", f"q{i}", "Q", *[None] * 5) for i in range(10)]
    with pytest.raises(XlsFormError, match="lignes"):
        parse_xlsform(_book(rows))
    monkeypatch.setattr(parser, "MAX_SHEETS", 2)
    with pytest.raises(XlsFormError, match="feuilles"):
        parse_xlsform(_book([HEAD, ("text", "q", "Q", *[None] * 5)], extra_sheets=3))
    monkeypatch.setattr(parser, "MAX_CELL", 5)
    with pytest.raises(XlsFormError, match="trop longue"):
        parse_xlsform(_book([HEAD, ("text", "q", "x" * 50, *[None] * 5)]))


def test_decompressed_size_bomb_is_refused(monkeypatch):
    monkeypatch.setattr(parser, "MAX_UNCOMPRESSED", 100)
    with pytest.raises(XlsFormError, match="décompressé"):
        parse_xlsform(_thirty())


def test_formulas_are_never_evaluated():
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(list(HEAD))
    ws.append(["text", "q", "=1+1", None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    plan = parse_xlsform(buf.getvalue())
    # data_only : une formule jamais calculée (pas de cache Excel) vaut vide, jamais 2
    assert plan.form_fields[0]["label"] in ("q", "=1+1")
    assert plan.form_fields[0]["label"] != "2"


def test_xxe_payload_is_refused_cleanly():
    good = _book([HEAD, ("text", "q", "Q", *[None] * 5)])
    src = zipfile.ZipFile(io.BytesIO(good))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith("sheet1.xml"):
                data = (
                    b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>'
                    + data
                )
            dst.writestr(item, data)
    # defusedxml (via openpyxl) refuse l'entité : erreur propre, jamais de contenu local lu
    with pytest.raises(XlsFormError, match="illisible"):
        parse_xlsform(out.getvalue())


def test_translated_label_choice_ignores_column_order():
    """AUD-115 : sans `label`, le libellé retenu ne dépend pas de l'ordre des colonnes."""
    from app.xlsform.parser import _label

    assert _label({"label::english": "Name", "label::français": "Nom"}) == "Nom"
    assert _label({"label::français": "Nom", "label::english": "Name"}) == "Nom"
    assert _label({"label::german": "Name", "label::english": "Name en"}) == "Name en"
    assert _label({"label": "Base", "label::français": "Nom"}) == "Base"


# ------------------------------------------------------------------ L8 T6 : fidélité XLSForm

SURVEY_HEAD = ("type", "name", "label", "required", "relevant", "default", "calculation")


def _plan(rows, choices=()):
    return parse_xlsform(_book([SURVEY_HEAD, *rows], choices=choices))


def _rule(plan, rid):
    return next(r for r in plan.attribute_rules if r["id"] == rid)


def test_choice_labels_prefer_fr_then_first_translation_and_are_optional():
    plan = _plan(
        [("select_one l", "x", "X", None, None, None, None)],
        choices=[
            ("list_name", "name", "label::en", "label::fr"),
            ("l", "a", "Alpha", "Alfa"),
            ("l", "b", "Beta", None),
        ],
    )
    assert plan.field_domains["x"]["labels"] == {"a": "Alfa", "b": "Beta"}
    plain = _plan(
        [("select_one l", "x", "X", None, None, None, None)],
        choices=[("list_name", "name"), ("l", "a"), ("l", "b")],
    )
    assert "labels" not in plain.field_domains["x"]


def test_required_becomes_a_server_constraint_per_column_type():
    plan = _plan(
        [
            ("text", "nom", "Nom", "yes", None, None, None),
            ("integer", "n", "N", "yes", None, None, None),
            ("select_multiple l", "m", "M", "yes", None, None, None),
            ("text", "libre", "Libre", None, None, None, None),
        ],
        choices=[("list_name", "name"), ("l", "a")],
    )
    assert _rule(plan, "xf-required-nom")["expression"] == "nom != null && nom != ''"
    assert _rule(plan, "xf-required-n")["expression"] == "n != null"
    assert _rule(plan, "xf-required-m")["expression"] == "m != null && size(m) > 0"
    assert not [r for r in plan.attribute_rules if r["id"].endswith("libre")]


def test_required_under_relevant_is_only_enforced_when_the_field_is_visible():
    plan = _plan(
        [
            ("text", "a", "A", None, None, None, None),
            ("text", "b", "B", "yes", "${a} = 'x'", None, None),
        ]
    )
    assert _rule(plan, "xf-required-b")["expression"] == "!(a == 'x') || (b != null && b != '')"


def test_required_with_untranslatable_relevant_is_not_enforced_and_is_reported():
    plan = _plan(
        [
            ("text", "a", "A", None, None, None, None),
            ("text", "b", "B", "yes", "today() > ${a}", None, None),
        ]
    )
    assert not [r for r in plan.attribute_rules if r["id"] == "xf-required-b"]
    assert any(r.name == "b" and "required" in r.message for r in plan.report)


def test_default_values_become_default_mode_rules():
    plan = _plan(
        [
            ("integer", "i", "I", None, None, "7", None),
            ("decimal", "d", "D", None, None, "2.5", None),
            ("select_multiple l", "m", "M", None, None, "a b", None),
            ("date", "j", "J", None, None, "today()", None),
        ],
        choices=[("list_name", "name"), ("l", "a"), ("l", "b")],
    )
    for field_, expr in {"i": "7", "d": "2.5", "m": "['a', 'b']"}.items():
        r = _rule(plan, f"xf-default-{field_}")
        assert (r["kind"], r["mode"], r["expression"]) == ("calculation", "default", expr)
    assert any(r.name == "j" and r.code == "not_ported" for r in plan.report)


def test_calculate_is_ported_as_a_derived_rule_with_an_inferred_column_type():
    plan = _plan(
        [
            ("decimal", "l", "L", None, None, None, None),
            ("decimal", "w", "W", None, None, None, None),
            ("calculate", "aire", "Aire", None, None, None, "${l} * ${w}"),
            ("calculate", "auj", "Auj", None, None, None, "today()"),
        ]
    )
    types = {c["name"]: c["sqlType"] for c in plan.columns}
    assert types["aire"] == "double precision"
    r = _rule(plan, "xf-calc-aire")
    assert (r["kind"], r["mode"], r["field"], r["expression"]) == (
        "calculation",
        "derived",
        "aire",
        "l * w",
    )
    assert "auj" not in types
    assert any(r.name == "auj" and r.code == "not_ported" for r in plan.report)


def test_geometry_required_is_a_plan_flag():
    assert _plan([("geopoint", "g", "G", "yes", None, None, None)]).geometry_required is True
    assert _plan([("geopoint", "g", "G", None, None, None, None)]).geometry_required is False


def test_media_kinds_are_typed():
    plan = _plan(
        [
            ("image", "p", "P", None, None, None, None),
            ("audio", "a", "A", None, None, None, None),
            ("video", "v", "V", None, None, None, None),
            ("file", "f", "F", None, None, None, None),
        ]
    )
    assert {f["key"]: f["kind"] for f in plan.attachment_fields} == {
        "p": "image",
        "a": "audio",
        "v": "video",
        "f": "file",
    }
    # le champ du formulaire porte la même nature (pilote accept/capture côté shell)
    assert {f["name"]: f["kind"] for f in plan.form_fields} == {
        "p": "image",
        "a": "audio",
        "v": "video",
        "f": "file",
    }


# Parité : valeur de référence XPath (calculée à la main) contre l'évaluation CEL réelle.
_PARITY = [
    ("if(${n} > 3, 'grand', 'petit')", {"n": 5, "s": "abc"}, "grand"),
    ("if(${n} > 3, 'grand', 'petit')", {"n": 1, "s": "abc"}, "petit"),
    ("string-length(${s})", {"n": 1, "s": "abcd"}, 4),
    ("string-length(${s})", {"n": 1, "s": ""}, 0),
    ("coalesce(${s}, 'defaut')", {"n": 1, "s": "x"}, "x"),
    ("coalesce(${s}, 'defaut')", {"n": 1, "s": None}, "defaut"),
    ("regex(${s}, '^[0-9]+$')", {"n": 1, "s": "123"}, True),
    ("regex(${s}, '^[0-9]+$')", {"n": 1, "s": "12a"}, False),
    ("regex(${s}, '[0-9]')", {"n": 1, "s": "ab1"}, True),  # XPath : non ancré
    ("${n} mod 3", {"n": 10, "s": ""}, 1),
    ("not(${n} = 1) and ${n} < 9", {"n": 5, "s": ""}, True),
]


@pytest.mark.parametrize(("xpath", "props", "expected"), _PARITY)
def test_to_cel_matches_xpath_reference_values(xpath, props, expected):
    from app.collections.introspection import ColumnInfo, TableInfo
    from app.features.cel_rules import evaluate

    names = {"n": "n", "s": "s", "t": "t"}
    info = TableInfo(
        table_name="x",
        pk_column="id",
        geometry_column=None,
        geometry_type=None,
        srid=None,
        columns=[ColumnInfo(name=k, type="string", required=False) for k in names],
    )
    assert evaluate(info, to_cel(xpath, names), props) == expected


def test_today_has_no_cel_equivalent_and_is_refused():
    with pytest.raises(Exception, match="today"):
        to_cel("today()", {})


def test_translations_kept_in_form_field():
    """AUD-115 c : `label::en` / `hint::en` conservés ; langue hors fr/en ignorée."""
    plan = parse_xlsform(
        _book(
            [
                ("type", "name", "label", "label::English (en)", "label::de", "hint::en"),
                ("text", "nom", "Nom", "Name", "Nachname", "Your name"),
            ]
        )
    )
    field = plan.form_fields[0]
    assert field["label"] == "Nom"
    assert field["labels"] == {"en": "Name"}
    assert field["hints"] == {"en": "Your name"}
    assert any(r.code == "ignored" and "label::de" in r.message for r in plan.report)
